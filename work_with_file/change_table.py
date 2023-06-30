from docx import Document
# выравнивание тексиа
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
def replace_tags_in_table(Report_name, table_tags):
    doc = Document(Report_name)
    table = doc.tables[-1]
    # председатель
    row = table.rows[0].cells
    row[-1].text = table_tags[0]

    # члены ГЭК
    for i, name in enumerate(table_tags[1]):
        table.add_row()
        table.cell(-1, 0).merge(table.cell(-1, 1))
        row = table.rows[-1].cells
        row[-1].text = str(i+1) + ". " + name

    # Руководитель
    table.add_row()
    row = table.rows[-1].cells
    row[0].text = "Руководитель"
    row[-1].text = table_tags[2]

    # Консультант, если есть
    if table_tags[3]:
        table.add_row()
        row = table.rows[-1].cells
        row[0].text = "Консультант"
        row[-1].text = table_tags[3]

    doc.save(Report_name)

#replace_tags_in_table("../www.docx")
